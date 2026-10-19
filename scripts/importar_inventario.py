"""Importa el inventario de Ruedas Darío (Excel) a src/data/catalogo.json.

Uso:  python3 scripts/importar_inventario.py RUTA_AL_EXCEL.xlsx [--dry]

- Agrupa las referencias (SKU) del Excel en productos "Producto – Marca" según la
  tabla REGLAS (la primera regla que coincide gana).
- Los productos escritos a mano que ya están en el catálogo (ids de MANUALES) se
  conservan tal cual; solo se les suma la compatibilidad detectada.
- La compatibilidad se deduce de los vehículos que mencionan los nombres de las
  referencias (tabla VEHICULOS). No se publica stock, códigos ni cantidades.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CATALOGO = ROOT / 'src/data/catalogo.json'

# ---------------------------------------------------------------- lectura

def u(s):
    s = unicodedata.normalize('NFD', str(s or '').upper()).encode('ascii', 'ignore').decode()
    return ' '.join(s.split())


def slug(t):
    t = unicodedata.normalize('NFD', t.lower()).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', t).strip('-')


ALIAS = {
    'WEATHER TECH': 'WEATHERTECH', 'POWER STOP': 'POWERSTOP', 'HAWK PERFORMANCE': 'HAWK',
    'OME': 'OLD MAN EMU', 'TOYO TIRES': 'TOYO', 'BF GOODRICH': 'BFGOODRICH', 'TRIPLE 555': '555',
    'MT64': 'MT-64', 'BP51': 'BP-51', 'FULGOR BLACK EDITION': 'FULGOR', 'APLUS A610': 'APLUS',
    'HILO XU1': 'HILO', 'PNEUS VANTAGE': 'VANTAGE', "WOLF'S HEAD": 'WOLF', 'VPR': 'VIPER',
    'ROLL N ROCK': 'ROLL-N-LOCK',
}
# Marcas que se buscan en el nombre cuando el Excel no trae la columna Marca.
MARCAS_EXTRA = [
    'K&N', 'ARB', 'VIPER', 'VPR', 'KATD', 'WEATHERTECH', 'WEATHER TECH', 'PROMAT', 'TRD', 'JBL',
    'OPTIMA', 'SHELL', 'MOTUL', 'MOTORCRAFT', "WOLF'S HEAD", 'WOLF', 'BILSTEIN', 'MONROE', 'FABTECH',
    'ROUGH COUNTRY', 'CAMBURG', 'HAWK', 'WAGNER', 'POWERSTOP', 'POWER STOP', 'MOOG', '555', 'OME',
    'OLD MAN EMU', 'KSP', 'VP RACING', 'SCEPTER', 'DURABOX', 'PIRELLI', 'TOYOTA', 'KUMHO', 'NITTO',
    'TOYO', 'BF GOODRICH', 'BFGOODRICH', 'FUEL', 'BLACK RHINO', 'BWA', 'NICHE', 'NOCO', 'KENWOOD',
    'DS18', 'INFINITY', 'MAXXIS', 'FALKEN', 'MICHELIN', 'DUNLOP', 'BRIDGESTONE', 'YOKOHAMA', 'GENERAL',
    'CONTINENTAL', 'MT-64', 'MT64', 'BP-51', 'BP51', 'KYB', 'FULGOR', 'DURACOVER', 'COVERLOCK',
    'VITOUR', 'FUJISAKI', 'MAXTREK', 'AEROKLAS', 'RUGGED LINER', 'RIGID', 'VISION', 'BRAVA',
    'QUALID', 'MILLARD', 'AMSOIL', 'MAXILUB', 'STP', 'AISIN', 'METALCAR', 'INGOTRIA', 'BENEXT',
    'KEKO', 'CURT', 'TMT', 'WARN', 'FACTOR55', 'DAMON', 'PROPARTS', 'ROLL N ROCK', 'PNEUS VANTAGE',
    'VANTAGE', 'FIREMAX', 'TOURADOR', 'TRIANGLE', 'ANNAITE', 'APLUS', 'ARIVO', 'COMPASAL',
    'JOYROAD', 'LIONLORD', 'ROYAL BLACK', 'RAPID', 'ILINK', 'BLACK EAGLE', 'THUNOLER PLUS',
]


def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    skus, marca_por_codigo = {}, {}
    for hoja in ('G1 sin ventas', 'G2 mas de 5 meses'):
        for r in wb[hoja].iter_rows(min_row=10, values_only=True):
            if r[0] and isinstance(r[9], int):
                marca = u(r[2])
                skus[r[9]] = {'cat': u(r[3]), 'nombre': str(r[1]).strip(), 'marca': marca}
                if marca:
                    marca_por_codigo[r[0]] = marca
    for r in wb['Reposicion SKU'].iter_rows(min_row=9, values_only=True):
        if r[0] and isinstance(r[11], int) and r[11] not in skus:
            skus[r[11]] = {'cat': u(r[0]), 'nombre': str(r[2]).strip(), 'marca': marca_por_codigo.get(r[1], '')}

    conocidas = {s['marca'] for s in skus.values()} | {u(m) for m in MARCAS_EXTRA}
    conocidas -= {'', 'SIN MARCA'}
    conocidas = sorted(conocidas, key=len, reverse=True)
    for s in skus.values():
        s['n'] = u(s['nombre'])
        m = s['marca'] if s['marca'] not in ('', 'SIN MARCA') else ''
        if not m:
            texto = ' ' + re.sub(r'[\[\]()/,]', ' ', s['n']) + ' '
            m = next((b for b in conocidas if re.search(r'(?<![A-Z0-9])' + re.escape(b) + r'(?![A-Z0-9])', texto)), '')
        s['b'] = ALIAS.get(m, m)
    return list(skus.values())

# ---------------------------------------------------------------- vehículos

# (patrón sobre el nombre normalizado, etiqueta). Las etiquetas contienen los
# términos de VEHICULOS en src/config.ts para que funcione el filtro.
VEHICULOS = [
    (r'4\s?-?RUNNER|\b4R\b', 'Toyota 4Runner'),
    (r'\bREVO\b', 'Toyota Hilux Revo'),
    (r'KAVAK|KAWAK|\bKAVK\b|\bVIGO\b', 'Toyota Hilux Kavak'),
    (r'HILUX', 'Toyota Hilux'),
    (r'DUBAI', 'Toyota Fortuner Dubai'),
    (r'FORTUNER|\bFORT\b', 'Toyota Fortuner'),
    (r'\bLC\s?300\b', 'Toyota Land Cruiser 300'),
    (r'\bLC\s?250\b|PRADO 250', 'Toyota Land Cruiser 250'),
    (r'\bLC\s?200\b', 'Toyota Land Cruiser 200'),
    (r'MACHITO|MACHT|SERIE ?7[0-9]|\bS7[0-9]\b', 'Toyota Land Cruiser serie 70 (Machito)'),
    (r'HEMBRITA', 'Toyota Land Cruiser Hembrita'),
    (r'LAND ?CRU|\bLANDCR', 'Toyota Land Cruiser'),
    (r'AUTANA|\bAUTA\b', 'Toyota Autana (serie 80)'),
    (r'RORAIMA|\bRORA\b', 'Toyota Roraima'),
    (r'\bPRADO\b|\bPRAD\b', 'Toyota Prado'),
    (r'\bFJ\b', 'Toyota FJ Cruiser'),
    (r'SEQUOIA|\bSEQ\b', 'Toyota Sequoia'),
    (r'TACOMA|\bTAC\b', 'Toyota Tacoma'),
    (r'TUNDRA|\bTUND\b', 'Toyota Tundra'),
    (r'COROLLA|\bCOROL\b', 'Toyota Corolla'),
    (r'YARIS', 'Toyota Yaris'),
    (r'CAMRY', 'Toyota Camry'),
    (r'RANGER', 'Ford Ranger'),
    (r'\bF-?150\b|BRONCO', 'Ford F150'),
    (r'\bF-?2[57]0\b|\bF-?350\b|SUPER DUTY', 'Ford Super Duty'),
    (r'EXPLORER|\bEXPLO\b', 'Ford Explorer'),
    (r'EXPEDITION', 'Ford Expedition'),
    (r'ESCAPE', 'Ford Escape'),
    (r'TERRITORY', 'Ford Territory'),
    (r'EVEREST', 'Ford Everest'),
    (r'SILVERADO|CHEYENNE', 'Chevrolet Silverado'),
    (r'TAHOE', 'Chevrolet Tahoe'),
    (r'AVALANCH', 'Chevrolet Avalanche'),
    (r'SUBURBAN', 'Chevrolet Suburban'),
    (r'\bWK2\b|4G WK', 'Jeep Grand Cherokee WK2'),
    (r'\bWK\b', 'Jeep Grand Cherokee WK'),
    (r'LIBERTY|\bKK\b', 'Jeep Liberty KK'),
    (r'\bL-?200\b|TRITON', 'Mitsubishi L200'),
    (r'MONTERO', 'Mitsubishi Montero'),
    (r'NAVARA|NP300', 'Nissan Navara'),
    (r'FRONTIER', 'Nissan Frontier'),
    (r'PATHFINDER', 'Nissan Pathfinder'),
    (r'\bD-?MAX\b', 'Chevrolet D-Max'),
    (r'JAC T9|\bT9\b', 'JAC T9'),
    (r'KAICENE|\bF-?70\b', 'Changan Kaicene F70'),
]
VEH_RE = [(re.compile(p), label) for p, label in VEHICULOS]
GENERICAS = {'Toyota Hilux': ('Toyota Hilux Revo', 'Toyota Hilux Kavak'),
             'Toyota Fortuner': ('Toyota Fortuner Dubai',),
             'Toyota Land Cruiser': tuple(l for _, l in VEHICULOS if l.startswith('Toyota Land Cruiser '))}


def vehiculos(n):
    found = [label for rx, label in VEH_RE if rx.search(n)]
    # "Hilux" a secas sobra si ya se detectó Revo/Kavak, etc.
    return [l for l in found if not any(e in found for e in GENERICAS.get(l, ()))]

# ---------------------------------------------------------------- reglas

def has(s, *kws):
    return any(k in s['n'] for k in kws)


# Productos escritos a mano (los 40 originales): se conservan nombre, textos y
# compatibilidad; el script solo les suma vehículos. El resto se regenera.
MANUALES = {
    'nitro-charger-sport', 'mt64', 'bp51', 'rough-country', 'fabtech', 'bilstein-b6', 'monroe',
    'camburg', 'hawk', 'wagner', 'powerstop', 'moog-tren', '555-tren', 'moog-rodamientos',
    'optima-yellow', 'optima-red', 'optima-blue', 'aceite-shell', 'aceite-motorcraft',
    'aceite-motul', 'aceite-wolf', 'kn-filtros', 'trd-intake', 'modulos-gasolina', 'bedliner',
    'santa-maria', 'tapa-trifold', 'lonas', 'estribos-viper', 'estribos-trd', 'skid-arb',
    'defensa-viper', 'barras-tiro', 'faros-fortuner', 'luces-katd', 'alfombras-katd',
    'alfombras-weathertech', 'alfombras-promat', 'volantes', 'jbl',
}

CATEGORIAS = [
    ('cauchos', 'Cauchos', 'Cauchos para camionetas, autos y carga en muchas medidas'),
    ('rines', 'Rines', 'Rines, separadores y centralizadores'),
    ('suspension', 'Suspensión y amortiguación', 'Kits de suspensión, amortiguadores, espirales, ballestas y tijeras'),
    ('frenos', 'Frenos y tren delantero', 'Pastillas, discos, tren delantero y rodamientos'),
    ('motor', 'Motor, aceites y baterías', 'Aceites, filtros, intake, baterías y cargadores'),
    ('exterior', 'Exterior y protección', 'Defensas, estribos, bedliner, tapas, lonas y luces'),
    ('interior', 'Interior y sonido', 'Alfombras, volantes, pantallas y sonido'),
    ('accesorios', 'Accesorios y herramientas', 'Rescate, sensores TPMS, tuercas, merchandising y más'),
]

# (id, categoría del sitio, nombre, marca, detalle, condición). nombre=None → producto manual existente.
R = []


def regla(id, cat, nombre, marca, cond, detalle=None):
    R.append(dict(id=id, cat=cat, nombre=nombre, marca=marca, detalle=detalle, cond=cond))


def manual(id, cond):
    R.append(dict(id=id, cat=None, nombre=None, marca=None, detalle=None, cond=cond))


C = lambda *cats: (lambda s: s['cat'] in cats)

regla(None, None, None, None, C('MISCELANEOS', 'CATEGORIA'))  # golosinas y encabezados: fuera

# --- kits de suspensión
manual('bp51', C('BP-51'))
regla('bp51-medios-kits', 'suspension', 'Medios kits BP-51', 'Old Man Emu', C('MEDIOS KITS BP-51'), '4 amortiguadores y fitting kits')
manual('mt64', C('MT-64'))
regla('mt64-medios-kits', 'suspension', 'Medios kits MT64', 'Old Man Emu', C('MEDIOS KITS MT-64'), '4 amortiguadores y espirales')

# --- reglas por marca que cruzan categorías
manual('rough-country', lambda s: s['b'] == 'ROUGH COUNTRY')
manual('kn-filtros', lambda s: (s['b'] == 'K&N' or has(s, 'K&N')) and has(s, 'FILTRO', 'INTAKE'))
regla('kn-limpiaparabrisas', 'exterior', 'Limpiaparabrisas', 'K&N', lambda s: has(s, 'LIMPIAPARABRISAS'))
regla('kn-limpiador-combustible', 'motor', 'Limpiador de sistema de combustible', 'K&N', lambda s: has(s, 'FUEL SYSTEM'))
regla('noco', 'motor', 'Cargadores, arrancadores y accesorios', 'NOCO', lambda s: s['b'] == 'NOCO')
regla('compresor-arb', 'accesorios', 'Compresor Heavy Duty', 'ARB', lambda s: has(s, 'COMPRESOR'))
regla('separadores-ksp', 'rines', 'Separadores de rines', 'KSP', lambda s: has(s, 'SEPARADOR'))
regla('centralizadores-rines', 'rines', 'Centralizadores de rines', None, lambda s: has(s, 'CENTRALIZADOR') and has(s, 'RINES'))
regla('rines-otros', 'rines', 'Rines (varias marcas)', None, lambda s: s['cat'] == 'REPUESTOS' and re.match(r'\d{2}X\d', s['n']))  # rines cargados como repuesto
regla('leveling-ksp', 'suspension', 'Leveling kits', 'KSP', lambda s: has(s, 'LEVELING') and s['b'] == 'KSP')
regla('maxilub', 'motor', 'Aceite de transmisión, refrigerante y liga de frenos', 'Maxilub', lambda s: has(s, 'MAXILUB'))
manual('trd-intake', lambda s: has(s, 'INTAKE'))
manual('volantes', lambda s: has(s, 'VOLANTE'))
manual('modulos-gasolina', lambda s: has(s, 'MODULO') and has(s, 'GASOLINA'))

# --- Filtros (antes que Repuestos/Accesorios)
regla('filtros-motorcraft', 'motor', 'Filtros de aceite', 'Motorcraft', lambda s: s['b'] == 'MOTORCRAFT' and has(s, 'FILTRO'))
regla('refrigerante-qualid', 'motor', 'Refrigerante', 'Qualid', lambda s: has(s, 'QUALID') and has(s, 'REFRIGERANTE'))
regla('filtros-qualid', 'motor', 'Filtros de aceite y de aire', 'Qualid', lambda s: has(s, 'QUALID'))
regla('filtros-millard', 'motor', 'Filtros de aceite', 'Millard', lambda s: has(s, 'MILLARD'))
regla('filtros-toyota', 'motor', 'Filtros originales', 'Toyota', lambda s: has(s, 'FILTRO') and s['b'] == 'TOYOTA')
regla('filtros-otros', 'motor', 'Filtros de aceite y de cabina', None, lambda s: has(s, 'FILTRO'))

# --- Old Man Emu
OME = C('OLD MAN EMU')
regla('ome-repuestos-bp51', 'suspension', 'Repuestos BP-51', 'Old Man Emu', lambda s: OME(s) and has(s, 'BP-51', 'BP51', 'BP5'), 'Amortiguadores, bujes y fitting kits')
regla('ome-repuestos-mt64', 'suspension', 'Repuestos MT64', 'Old Man Emu', lambda s: OME(s) and has(s, 'MT64', 'MT-64'), 'Amortiguadores, espirales y bujes')
regla('ome-airbag', 'suspension', 'Kits de airbag para ballestas', 'Old Man Emu', lambda s: OME(s) and has(s, 'AIR BAG'))
regla('ome-gemelos', 'suspension', 'Gemelos engrasables', 'Old Man Emu', lambda s: OME(s) and has(s, 'GEMELO'))
regla('ome-amortiguadores', 'suspension', 'Amortiguadores', 'Old Man Emu', lambda s: OME(s) and has(s, 'AMORT'))
regla('ome-espirales', 'suspension', 'Espirales', 'Old Man Emu', lambda s: OME(s) and has(s, 'ESPIRAL'))
regla('ome-ballestas', 'suspension', 'Ballestas', 'Old Man Emu', lambda s: OME(s) and has(s, 'BALLESTA', 'LEAF SPRING'))
regla('ome-barra-estabilizadora', 'suspension', 'Kits para barra estabilizadora', 'Old Man Emu', lambda s: OME(s) and has(s, 'ESTAB'))
regla('ome-espaciadores', 'suspension', 'Espaciadores', 'Old Man Emu', lambda s: OME(s) and has(s, 'ESPACIADOR'))
regla('ome-mesetas', 'suspension', 'Mesetas y muñones', 'Old Man Emu', lambda s: OME(s) and has(s, 'MESETA', 'MUNON'))
regla('ome-bujes', 'suspension', 'Bujes y kits de adaptación', 'Old Man Emu', lambda s: OME(s) and has(s, 'BUJE', 'KIT PARA', 'FITTING'))
regla('ome-repuestos', 'suspension', 'Tornillos U y repuestos', 'Old Man Emu', OME)

# --- Moog
manual('moog-rodamientos', lambda s: s['cat'] == 'MOOG' and has(s, 'ROLINERA', 'RODAMIENTO', 'MOZO'))
manual('moog-tren', lambda s: s['cat'] == 'MOOG' or (s['b'] == 'MOOG' and s['cat'] == 'REPUESTOS'))

# --- Frenos
F = C('FRENOS')
manual('hawk', lambda s: F(s) and s['b'] == 'HAWK')
manual('powerstop', lambda s: F(s) and s['b'] == 'POWERSTOP')
manual('wagner', lambda s: F(s) and s['b'] == 'WAGNER')
regla('frenos-toyota', 'frenos', 'Pastillas y discos de freno originales', 'Toyota', lambda s: F(s) and s['b'] == 'TOYOTA')
regla('bandas-freno', 'frenos', 'Bandas de freno', None, lambda s: F(s) and has(s, 'BANDA'), 'Originales y Aisin')
regla('liga-frenos', 'frenos', 'Liga de frenos', None, lambda s: F(s) and has(s, 'LIGA'))
regla('grasa-frenos', 'frenos', 'Grasa para frenos', 'Lubrisil', lambda s: F(s) and has(s, 'LUBRISIL', 'GRASA'))
regla('frenos-otras', 'frenos', 'Pastillas y discos de freno (otras marcas)', None, F)

# --- Repuestos y suspensión
RP = C('REPUESTOS')
manual('monroe', lambda s: s['b'] == 'MONROE')
manual('camburg', lambda s: s['b'] == 'CAMBURG')
manual('555-tren', lambda s: s['b'] == '555')
manual('bilstein-b6', lambda s: s['b'] == 'BILSTEIN')
manual('fabtech', lambda s: s['b'] == 'FABTECH')
regla('amortiguadores-kyb', 'suspension', 'Amortiguadores', 'KYB', lambda s: s['b'] == 'KYB')
regla('repuestos-toyota', 'frenos', 'Repuestos originales', 'Toyota', lambda s: RP(s) and s['b'] == 'TOYOTA', 'Terminales, rótulas, rodamientos, estoperas y más')
regla('espirales-metalcar', 'suspension', 'Espirales', 'Metalcar', lambda s: s['b'] == 'METALCAR')
regla('guardapolvos', 'suspension', 'Guardapolvos de amortiguador', None, lambda s: RP(s) and has(s, 'GUARDA POLVO', 'GUARDAPOLVO'))
regla('corrector-diferencial', 'suspension', 'Corrector de diferencial', 'TMT', lambda s: has(s, 'DIFF DROP', 'CORRECTOR DIFERENCIAL'))
regla('rep-amortiguadores', 'suspension', 'Amortiguadores (otras marcas)', None, lambda s: RP(s) and has(s, 'AMORT'))
regla('rep-tijeras', 'suspension', 'Tijeras', None, lambda s: RP(s) and has(s, 'TIJERA'))
regla('rep-estabilizadora', 'frenos', 'Barras estabilizadoras y gomas', None, lambda s: RP(s) and has(s, 'ESTAB'))
regla('rep-rodamientos', 'frenos', 'Rodamientos y rolineras', None, lambda s: RP(s) and has(s, 'RODAMIENTO', 'ROLINERA'))
regla('rep-rotulas', 'frenos', 'Rótulas, terminales y mesetas', None, lambda s: RP(s) and has(s, 'ROTULA', 'TERMINAL', 'MUNON', 'MESETA'))
regla('rep-bujes', 'frenos', 'Bujes de barra y suspensión', None, lambda s: RP(s) and has(s, 'BUJE'))
regla('rep-estoperas', 'motor', 'Estoperas y topes', None, lambda s: RP(s) and has(s, 'ESTOPERA', 'TOPE'))
regla('rep-otros', 'frenos', 'Otros repuestos', None, RP, 'Gomas de trípode, grasas, anillos y más')
regla('suspension-otros', 'suspension', 'Espirales y kits de nivelación', None, C('SUSPENSION'))

# --- Aceites y filtros
A = C('ACEITES')
manual('aceite-shell', lambda s: s['b'] == 'SHELL')
manual('aceite-motorcraft', lambda s: s['b'] == 'MOTORCRAFT' and not has(s, 'FILTRO'))
manual('aceite-motul', lambda s: s['b'] == 'MOTUL')
manual('aceite-wolf', lambda s: s['b'] == 'WOLF')
regla('aceite-brava', 'motor', 'Aceites', 'Brava', lambda s: s['b'] == 'BRAVA')
regla('aceite-kn', 'motor', 'Aceites', 'K&N', lambda s: A(s) and s['b'] == 'K&N')
regla('aceite-toyota', 'motor', 'Aceites', 'Toyota', lambda s: A(s) and s['b'] == 'TOYOTA')
regla('aceite-thunoler', 'motor', 'Aceites', 'Thunoler Plus', lambda s: has(s, 'THUNOLER'))
regla('amsoil', 'motor', 'Aceites, grasas y líquido de frenos', 'Amsoil', lambda s: A(s) and has(s, 'AMSOIL', 'CJ '))
regla('aditivo-stp', 'motor', 'Aditivo', 'STP', lambda s: s['b'] == 'STP')
regla('aceites-otros', 'motor', 'Aceites (otras marcas)', None, A)

# --- Baterías
B = C('BATERIAS')
manual('optima-yellow', lambda s: B(s) and s['b'] == 'OPTIMA' and has(s, 'AMARILLA', 'YELLOW'))
manual('optima-red', lambda s: B(s) and s['b'] == 'OPTIMA' and has(s, 'ROJA', 'RED'))
manual('optima-blue', lambda s: B(s) and s['b'] == 'OPTIMA')
regla('baterias-fulgor', 'motor', 'Baterías', 'Fulgor', lambda s: B(s) and s['b'] == 'FULGOR')
regla('baterias-otras', 'motor', 'Baterías (otras marcas)', None, B, 'Bestia Negra y otras')

# --- Bedliner, luces, sonido
manual('bedliner', C('BEDLINER'))
manual('luces-katd', lambda s: s['cat'] == 'LUCES' and s['b'] == 'KATD')
regla('faros-arb', 'exterior', 'Faros Intensity IQ', 'ARB', lambda s: s['cat'] == 'LUCES' and s['b'] == 'ARB')
regla('luces-rigid', 'exterior', 'Faros y barras LED', 'Rigid', lambda s: s['b'] == 'RIGID')
regla('luces-viper', 'exterior', 'Barras y faros LED', 'Viper', lambda s: s['b'] == 'VIPER' and has(s, 'BARRA LED', 'FAROS LED'))
regla('luces-vision', 'exterior', 'Bombillos LED', 'Vision', lambda s: s['b'] == 'VISION')
regla('luces-otras', 'exterior', 'Estrobos y luces auxiliares', None, lambda s: s['cat'] == 'LUCES' or (s['cat'] == 'ACCESORIOS' and has(s, 'LUCES', 'ESTROBO')))
manual('jbl', lambda s: s['b'] == 'JBL' or has(s, 'FLIP 6', 'CHARGE 5'))
regla('kenwood', 'interior', 'Pantallas, amplificadores y cornetas', 'Kenwood', lambda s: s['b'] == 'KENWOOD')
regla('ds18', 'interior', 'Bajos, medios y accesorios de sonido', 'DS18', lambda s: has(s, 'DS18'))
regla('infinity', 'interior', 'Bajos', 'Infinity', lambda s: has(s, 'INFINITY'))
regla('sonido-otros', 'interior', 'Amplificadores y bajos (otras marcas)', None, C('SONIDO'))

# --- Alfombras
AL = C('ALFOMBRAS')
manual('alfombras-katd', lambda s: AL(s) and s['b'] == 'KATD')
manual('alfombras-weathertech', lambda s: AL(s) and (s['b'] == 'WEATHERTECH' or has(s, 'CARGOLINER PRADO')))
manual('alfombras-promat', lambda s: AL(s) and s['b'] == 'PROMAT')
regla('alfombras-arb', 'interior', 'Alfombras y cargoliners', 'ARB', lambda s: AL(s) and s['b'] == 'ARB')
regla('alfombras-medida', 'interior', 'Alfombras a medida por vehículo', None, AL)

# --- Rines
RI = C('RINES')
for marca in ('BLACK RHINO', 'BWA', 'FUEL', 'NICHE', 'TRD'):
    regla(f'rines-{slug(marca)}', 'rines', 'Rines', marca.title() if marca != 'TRD' else 'TRD',
          (lambda m: lambda s: RI(s) and s['b'] == m)(marca))
regla('rines-toyota', 'rines', 'Rines originales', 'Toyota', lambda s: RI(s) and s['b'] == 'TOYOTA')
regla('rines-otros', 'rines', 'Rines (varias marcas)', None, RI)

# --- Cauchos: un producto por marca (se generan abajo)
CAUCHO_MARCA = {'BFGOODRICH': 'BFGoodrich', 'ILINK': 'iLink', 'CST': 'CST', 'V-NETIX': 'V-Netix'}

# --- Accesorios (exterior / interior / accesorios)
AC = C('ACCESORIOS')
manual('santa-maria', lambda s: has(s, 'SANTA MARIA'))
manual('tapa-trifold', lambda s: has(s, 'TAPA DURA', 'TAPA UNIVERSAL'))
manual('lonas', lambda s: has(s, 'LONA'))
manual('barras-tiro', lambda s: has(s, 'BARRA DE TIRO', 'TIRO SENCILLO', 'BALL MOUNT'))
manual('estribos-trd', lambda s: has(s, 'ESTRIBO') and s['b'] == 'TRD')
manual('estribos-viper', lambda s: has(s, 'ESTRIBO') and s['b'] == 'VIPER')
regla('estribos-otros', 'exterior', 'Estribos (otras marcas)', None, lambda s: has(s, 'ESTRIBO'))
manual('skid-arb', lambda s: has(s, 'SKID') and s['b'] == 'ARB')
manual('defensa-viper', lambda s: s['b'] == 'VIPER' and has(s, 'SKID', 'DEFENSA', 'PETO', 'PARACHOQUE'))
regla('skid-toyota', 'exterior', 'Skid plate tipo TRD', 'Toyota', lambda s: has(s, 'SKID', 'PETO'))
regla('trd-accesorios', 'exterior', 'Emblemas, rack de techo y tercer stop', 'TRD', lambda s: AC(s) and s['b'] == 'TRD')
regla('pantallas-toyota', 'interior', 'Pantallas con cámara 360', 'Toyota', lambda s: has(s, 'PANTALLA'))
regla('camara-360', 'interior', 'Cámara 360', 'Damon', lambda s: has(s, 'CAMARA'))
regla('arb-rescate', 'accesorios', 'Recovery points y guantes', 'ARB', lambda s: s['b'] == 'ARB' and has(s, 'RECOVERY'))
regla('rescate-otros', 'accesorios', 'Equipos de rescate', None, lambda s: has(s, 'RECOVERY', 'ULTRAHOOK', 'WINCH'), 'Factor55, Warn y otros')
regla('arb-medidor', 'accesorios', 'Medidor de presión digital', 'ARB', lambda s: s['b'] == 'ARB' and has(s, 'MEDIDOR'))
regla('arb-cava', 'accesorios', 'Cava plegable', 'ARB', lambda s: has(s, 'CAVA'))
regla('arb-linterna', 'accesorios', 'Linterna LED recargable', 'ARB', lambda s: has(s, 'LINTERNA'))
regla('arb-merch', 'accesorios', 'Gorras, ropa y accesorios', 'ARB', lambda s: AC(s) and s['b'] == 'ARB')
regla('ome-merch', 'accesorios', 'Gorras y sillas', 'Old Man Emu', lambda s: AC(s) and s['b'] in ('OLD MAN EMU', 'BP-51'))
regla('llaveros-pirelli', 'accesorios', 'Llaveros', 'Pirelli', lambda s: has(s, 'LLAVERO') and s['b'] == 'PIRELLI')
regla('llaveros-toyota', 'accesorios', 'Llaveros', 'Toyota', lambda s: has(s, 'LLAVERO'))
regla('ganchos-aeroklas', 'exterior', 'Ganchos metálicos', 'Aeroklas', lambda s: AC(s) and s['b'] == 'AEROKLAS')
regla('tpms-toyota', 'accesorios', 'Sensores TPMS', 'Toyota', lambda s: has(s, 'TPMS') and s['b'] == 'TOYOTA')
regla('tpms-otros', 'accesorios', 'Sensores TPMS (otras marcas)', None, lambda s: has(s, 'TPMS', 'SENSOR'))
regla('tuercas-toyota', 'accesorios', 'Tuercas y espárragos originales', 'Toyota', lambda s: has(s, 'TUERCA', 'ESPARRAGO') and s['b'] == 'TOYOTA')
regla('tuercas', 'accesorios', 'Tuercas, seguros y llaves', None, lambda s: has(s, 'TUERCA', 'SEGURO', 'ESPARRAGO'))
regla('valvulas', 'accesorios', 'Válvulas', None, lambda s: has(s, 'VALVULA'))
regla('bidones-scepter', 'accesorios', 'Bidones', 'Scepter', lambda s: has(s, 'BIDON'))
regla('octane-vp', 'motor', 'Octane booster', 'VP Racing', lambda s: has(s, 'OCTANE'))
regla('porta-herramientas', 'accesorios', 'Porta herramientas', 'Durabox', lambda s: s['b'] == 'DURABOX')
regla('porta-equipaje', 'exterior', 'Porta equipaje', 'Proparts', lambda s: has(s, 'PORTA EQUIPAJE'))
regla('starlink-viper', 'accesorios', 'Base de antena Starlink Mini', 'Viper', lambda s: has(s, 'STARLINK'))
regla('space-arm', 'suspension', 'Barras estabilizadoras Space Arm', None, lambda s: has(s, 'SPACE ARM'))
regla('accesorios-otros', 'accesorios', 'Otros accesorios', None, AC)


def asignar(s):
    if s['cat'] == 'CAUCHOS':
        return 'caucho'
    for r in R:
        if r['cond'](s):
            return r
    return 'sin-regla'

# ---------------------------------------------------------------- main

def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        sys.exit(__doc__)
    dry = '--dry' in sys.argv
    skus = leer_excel(args[0])
    actual = json.loads(CATALOGO.read_text())
    manuales = {p['id']: (c['id'], p) for c in actual['categorias'] for p in c['productos'] if p['id'] in MANUALES}
    if set(manuales) != MANUALES:
        sys.exit(f'Faltan productos manuales en el catálogo: {sorted(MANUALES - set(manuales))}')

    defs = {r['id']: r for r in R if r['id']}
    for r in R:
        if r['nombre'] is None and r['id'] and r['id'] not in manuales:
            sys.exit(f'Regla manual sin producto en el catálogo: {r["id"]}')
    for r in R:
        if r['nombre'] is not None and r['id'] in manuales:
            sys.exit(f'Id nuevo repetido con uno manual: {r["id"]}')

    grupos, sin_regla, excluidos = {}, [], 0
    for s in skus:
        r = asignar(s)
        if r == 'sin-regla':
            sin_regla.append(s); continue
        if r == 'caucho':
            marca = s['b'] or ''
            pid = f'cauchos-{slug(marca)}' if marca else 'cauchos-otras'
            if pid not in defs:
                defs[pid] = dict(id=pid, cat='cauchos', marca=CAUCHO_MARCA.get(marca, marca.title()) if marca else None,
                                 nombre='Cauchos' if marca else 'Cauchos (otras marcas)', detalle=None)
        elif r['id'] is None:
            excluidos += 1; continue
        else:
            pid = r['id']
        grupos.setdefault(pid, []).append(s)

    # Armar categorías: primero los manuales en su orden, luego los nuevos.
    salida = {cid: {'id': cid, 'nombre': nombre, 'subtitulo': sub, 'productos': []} for cid, nombre, sub in CATEGORIAS}
    for cid, p in manuales.values():
        compat = list(p['compatibilidad'])
        for s in grupos.get(p['id'], []):
            for v in vehiculos(s['n']):
                if not any(slug(v) in slug(c) or slug(c) in slug(v) for c in compat):
                    compat.append(v)
        salida[cid]['productos'].append({**p, 'compatibilidad': compat})

    orden_nuevos = [r['id'] for r in R if r['id'] and r['nombre']]
    cauchos = sorted((k for k in defs if k.startswith('cauchos-') and k != 'cauchos-otras'), key=lambda k: defs[k]['marca'])
    orden_nuevos = list(dict.fromkeys(cauchos + ['cauchos-otras'] + orden_nuevos))
    vacios = []
    for pid in orden_nuevos:
        if pid not in grupos:
            if pid in defs and defs[pid]['nombre'] and not pid.startswith('cauchos-'):
                vacios.append(pid)
            continue
        d = defs[pid]
        compat = []
        # Cauchos y rines van por medida: sirven para cualquier camioneta que la use.
        for s in grupos[pid] if d['cat'] not in ('cauchos', 'rines') else []:
            for v in vehiculos(s['n']):
                if v not in compat:
                    compat.append(v)
        p = {'id': pid, 'nombre': d['nombre']}
        if d['marca']:
            p['marca'] = d['marca']
        if d['detalle']:
            p['detalle'] = d['detalle']
        orden = [label for _, label in VEHICULOS]
        p['compatibilidad'] = sorted(compat, key=orden.index)
        salida[d['cat']]['productos'].append(p)

    total = sum(len(c['productos']) for c in salida.values())
    print(f'SKU leídos: {len(skus)} · excluidos: {excluidos} · sin regla: {len(sin_regla)}')
    for s in sin_regla[:30]:
        print('   SIN REGLA', s['cat'], '|', s['nombre'])
    print(f'Productos: {total} (manuales {len(manuales)}, nuevos {total - len(manuales)})')
    for c in salida.values():
        print(f'   {c["id"]:12} {len(c["productos"])}')
    sin_skus = [pid for pid in manuales if pid not in grupos]
    print('Manuales sin referencias en el Excel:', ', '.join(sin_skus) or '—')
    print('Reglas nuevas sin referencias:', ', '.join(vacios) or '—')
    if sin_regla:
        sys.exit('Hay referencias sin regla: no se escribe el catálogo.')
    if not dry:
        CATALOGO.write_text(json.dumps({'categorias': list(salida.values())}, ensure_ascii=False, indent=2) + '\n')
        print('Escrito', CATALOGO.relative_to(ROOT))


if __name__ == '__main__':
    main()
